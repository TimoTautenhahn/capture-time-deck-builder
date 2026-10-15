import argparse
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable, List, Optional

from PIL import Image, ImageEnhance, ImageFilter, ImageOps
from pptx import Presentation
from pptx.util import Inches

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except Exception:
    pass


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".heif"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm"}
CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080
MIME_DEFAULT = "video/mp4"


@dataclass
class MediaItem:
    path: Path
    kind: str  # image | video
    capture_time: datetime
    width: int
    height: int
    synthetic: bool = False


def list_media_recursively(folder: Path) -> List[Path]:
    files: List[Path] = []
    for root, _, names in os.walk(folder):
        for name in names:
            p = Path(root) / name
            ext = p.suffix.lower()
            if ext in IMAGE_EXTENSIONS or ext in VIDEO_EXTENSIONS:
                files.append(p)
    return files


def parse_exif_datetime(value: str) -> Optional[datetime]:
    if not value:
        return None
    for fmt in ("%Y:%m:%d %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def load_image(path: Path) -> Image.Image:
    with Image.open(path) as img:
        corrected = ImageOps.exif_transpose(img).convert("RGB")
    return corrected


def read_image_capture_time(path: Path) -> datetime:
    try:
        with Image.open(path) as img:
            exif = img.getexif()
            if exif:
                for tag in (36867, 36868, 306):
                    raw = exif.get(tag)
                    if raw:
                        parsed = parse_exif_datetime(str(raw))
                        if parsed:
                            return parsed
    except Exception:
        pass
    return datetime.fromtimestamp(path.stat().st_mtime)


def inspect_image(path: Path) -> MediaItem:
    img = load_image(path)
    width, height = img.size
    return MediaItem(path=path, kind="image", capture_time=read_image_capture_time(path), width=width, height=height)


def resolve_ffmpeg_executable() -> Optional[str]:
    system_ffmpeg = shutil.which("ffmpeg")
    if system_ffmpeg:
        return system_ffmpeg
    try:
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return None


def ffmpeg_probe_output(ffmpeg_exe: str, media_path: Path) -> str:
    proc = subprocess.run(
        [ffmpeg_exe, "-hide_banner", "-i", str(media_path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        check=False,
        encoding="utf-8",
        errors="ignore",
    )
    return proc.stderr or proc.stdout or ""


def parse_creation_time_from_ffmpeg_output(raw: str) -> Optional[datetime]:
    match = re.search(r"creation_time\s*:\s*([^\n\r]+)", raw)
    if not match:
        return None
    value = match.group(1).strip().replace("Z", "+00:00")
    for fmt in ("%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%d %H:%M:%S"):
        try:
            dt = datetime.strptime(value, fmt)
            if dt.tzinfo:
                return dt.astimezone().replace(tzinfo=None)
            return dt
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(value)
    except Exception:
        return None


def parse_video_dimensions_from_ffmpeg_output(raw: str) -> Optional[tuple[int, int]]:
    match = re.search(r"(\d{2,5})x(\d{2,5})", raw)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def inspect_video(path: Path, ffmpeg_exe: Optional[str]) -> MediaItem:
    capture_time = datetime.fromtimestamp(path.stat().st_mtime)
    width, height = 1920, 1080
    if ffmpeg_exe:
        raw = ffmpeg_probe_output(ffmpeg_exe, path)
        maybe_dt = parse_creation_time_from_ffmpeg_output(raw)
        maybe_size = parse_video_dimensions_from_ffmpeg_output(raw)
        if maybe_dt:
            capture_time = maybe_dt
        if maybe_size:
            width, height = maybe_size
    return MediaItem(path=path, kind="video", capture_time=capture_time, width=width, height=height)


def is_low_quality(item: MediaItem) -> bool:
    min_edge = min(item.width, item.height)
    area = item.width * item.height
    return min_edge < 900 or area < 1_000_000


def side_extend_background(img: Image.Image, canvas_size: tuple[int, int]) -> Image.Image:
    canvas_w, canvas_h = canvas_size
    resized = img.resize((max(1, int(img.width * canvas_h / img.height)), canvas_h), Image.Resampling.LANCZOS)
    strip_w = max(1, img.width // 10)
    left_strip = img.crop((0, 0, strip_w, img.height)).resize((canvas_w // 2, canvas_h), Image.Resampling.BICUBIC)
    right_strip = img.crop((img.width - strip_w, 0, img.width, img.height)).resize(
        (canvas_w - canvas_w // 2, canvas_h), Image.Resampling.BICUBIC
    )
    bg = Image.new("RGB", (canvas_w, canvas_h))
    bg.paste(left_strip, (0, 0))
    bg.paste(right_strip, (canvas_w // 2, 0))
    bg = bg.filter(ImageFilter.GaussianBlur(radius=45))
    bg = ImageEnhance.Brightness(bg).enhance(0.82)
    x = (canvas_w - resized.width) // 2
    bg.paste(resized, (x, 0))
    return bg


def blurred_full_background(img: Image.Image, canvas_size: tuple[int, int]) -> Image.Image:
    bg = ImageOps.fit(img, canvas_size, method=Image.Resampling.LANCZOS)
    bg = bg.filter(ImageFilter.GaussianBlur(radius=36))
    return ImageEnhance.Brightness(bg).enhance(0.78)


def render_stylized_canvas(img: Image.Image, width: int, height: int, low_quality: bool) -> Image.Image:
    canvas = Image.new("RGB", (CANVAS_WIDTH, CANVAS_HEIGHT), color=(20, 20, 20))
    canvas_ratio = CANVAS_WIDTH / CANVAS_HEIGHT
    image_ratio = width / height if height else 1.0

    if low_quality:
        bg = blurred_full_background(img, (CANVAS_WIDTH, CANVAS_HEIGHT))
        canvas.paste(bg, (0, 0))
        fg = ImageOps.contain(img, (int(CANVAS_WIDTH * 0.75), int(CANVAS_HEIGHT * 0.82)), method=Image.Resampling.LANCZOS)
        x = (CANVAS_WIDTH - fg.width) // 2
        y = (CANVAS_HEIGHT - fg.height) // 2
        canvas.paste(fg, (x, y))
        return canvas

    if image_ratio < canvas_ratio:
        bg = side_extend_background(img, (CANVAS_WIDTH, CANVAS_HEIGHT))
        canvas.paste(bg, (0, 0))
        fg = ImageOps.contain(img, (int(CANVAS_WIDTH * 0.94), CANVAS_HEIGHT), method=Image.Resampling.LANCZOS)
        x = (CANVAS_WIDTH - fg.width) // 2
        y = (CANVAS_HEIGHT - fg.height) // 2
        canvas.paste(fg, (x, y))
        return canvas

    fill = ImageOps.fit(img, (CANVAS_WIDTH, CANVAS_HEIGHT), method=Image.Resampling.LANCZOS)
    canvas.paste(fill, (0, 0))
    return canvas


def pil_to_png_bytes(image: Image.Image) -> io.BytesIO:
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    buf.seek(0)
    return buf


def extract_video_frame(video_path: Path, ffmpeg_exe: str, at_seconds: float = 0.7) -> Image.Image:
    cmd = [
        ffmpeg_exe,
        "-hide_banner",
        "-ss",
        f"{at_seconds}",
        "-i",
        str(video_path),
        "-frames:v",
        "1",
        "-f",
        "image2pipe",
        "-vcodec",
        "png",
        "-",
    ]
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    if proc.returncode != 0 or not proc.stdout:
        raise RuntimeError(f"Konnte kein Posterframe extrahieren: {video_path}")
    return Image.open(io.BytesIO(proc.stdout)).convert("RGB")


def ensure_mp4_for_embedding(video_path: Path, ffmpeg_exe: str, cache_dir: Path) -> Path:
    if video_path.suffix.lower() == ".mp4":
        return video_path
    cache_dir.mkdir(parents=True, exist_ok=True)
    target = cache_dir / f"{video_path.stem}.converted.mp4"
    if target.exists():
        return target
    cmd = [
        ffmpeg_exe,
        "-y",
        "-hide_banner",
        "-i",
        str(video_path),
        "-c:v",
        "libx264",
        "-preset",
        "medium",
        "-crf",
        "20",
        "-pix_fmt",
        "yuv420p",
        "-c:a",
        "aac",
        "-movflags",
        "+faststart",
        str(target),
    ]
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, check=False)
    if proc.returncode != 0:
        raise RuntimeError(f"Konvertierung fehlgeschlagen fuer {video_path}: {proc.stderr[:500]}")
    return target


def add_fullslide_image(prs: Presentation, rendered: Image.Image) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    stream = pil_to_png_bytes(rendered)
    slide.shapes.add_picture(stream, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)


def add_video_slide(prs: Presentation, video_path: Path, poster: Image.Image) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    poster_stream = pil_to_png_bytes(poster)
    mime = mimetypes.guess_type(str(video_path))[0] or MIME_DEFAULT
    slide.shapes.add_movie(
        str(video_path),
        Inches(0),
        Inches(0),
        width=prs.slide_width,
        height=prs.slide_height,
        mime_type=mime,
        poster_frame_image=poster_stream,
    )


def build_media_items(input_dir: Path, ffmpeg_exe: Optional[str]) -> List[MediaItem]:
    paths = list_media_recursively(input_dir)
    items: List[MediaItem] = []
    for p in paths:
        try:
            ext = p.suffix.lower()
            if ext in IMAGE_EXTENSIONS:
                items.append(inspect_image(p))
            else:
                items.append(inspect_video(p, ffmpeg_exe))
        except Exception as exc:
            print(f"Warnung: Datei wird uebersprungen ({p}): {exc}")
    items.sort(key=lambda x: (x.capture_time, str(x.path)))
    return items


def pick_test_items(all_items: List[MediaItem], output_dir: Path) -> List[MediaItem]:
    images = [i for i in all_items if i.kind == "image"]
    videos = [v for v in all_items if v.kind == "video"]
    if not images:
        raise RuntimeError("Keine Bilder gefunden. Mindestens ein Bild fuer den Test erforderlich.")

    portrait = [i for i in images if i.height > i.width]
    landscape = [i for i in images if i.width >= i.height]
    lowq = [i for i in images if is_low_quality(i)]
    large = sorted(images, key=lambda x: x.width * x.height, reverse=True)

    selected_images: List[MediaItem] = []
    used = set()
    for group in (portrait, landscape, lowq, large):
        candidate = next((x for x in group if x.path not in used), None)
        if candidate:
            selected_images.append(candidate)
            used.add(candidate.path)

    if len(selected_images) < 4:
        base = images[0]
        base_img = load_image(base.path)
        temp_dir = output_dir / "_autogen_test_assets"
        temp_dir.mkdir(parents=True, exist_ok=True)
        now = datetime.now().replace(microsecond=0)
        variants = [
            ("portrait.jpg", ImageOps.contain(base_img, (900, 1600), method=Image.Resampling.LANCZOS)),
            ("landscape.jpg", ImageOps.contain(base_img, (2200, 1200), method=Image.Resampling.LANCZOS)),
            ("tiny_lowq.jpg", ImageOps.contain(base_img, (360, 240), method=Image.Resampling.BILINEAR)),
            ("huge.jpg", ImageOps.contain(base_img, (5000, 3300), method=Image.Resampling.LANCZOS)),
        ]
        selected_images = []
        for idx, (name, variant) in enumerate(variants):
            file_path = temp_dir / name
            variant.save(file_path, quality=85)
            w, h = variant.size
            selected_images.append(
                MediaItem(
                    path=file_path,
                    kind="image",
                    capture_time=now + timedelta(seconds=idx),
                    width=w,
                    height=h,
                    synthetic=True,
                )
            )

    selected_images = sorted(selected_images[:4], key=lambda x: (x.capture_time, str(x.path)))
    if videos:
        first_video = sorted(videos, key=lambda x: (x.capture_time, str(x.path)))[0]
        return selected_images + [first_video]
    return selected_images


def create_presentation(
    media_items: Iterable[MediaItem],
    output_file: Path,
    ffmpeg_exe: Optional[str],
    conversion_cache_dir: Path,
) -> dict:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    stats = {"image_slides": 0, "video_poster_slides": 0, "video_embed_slides": 0, "media_items": 0}

    for item in media_items:
        stats["media_items"] += 1
        if item.kind == "image":
            img = load_image(item.path)
            rendered = render_stylized_canvas(img, item.width, item.height, is_low_quality(item))
            add_fullslide_image(prs, rendered)
            stats["image_slides"] += 1
            continue

        if not ffmpeg_exe:
            print(f"Warnung: Video ohne FFmpeg uebersprungen ({item.path})")
            continue

        try:
            poster_img = extract_video_frame(item.path, ffmpeg_exe)
            poster_canvas = render_stylized_canvas(
                poster_img, poster_img.width, poster_img.height, is_low_quality(item)
            )
            add_fullslide_image(prs, poster_canvas)
            stats["video_poster_slides"] += 1

            embeddable = ensure_mp4_for_embedding(item.path, ffmpeg_exe, conversion_cache_dir)
            add_video_slide(prs, embeddable, poster_canvas)
            stats["video_embed_slides"] += 1
        except Exception as exc:
            print(f"Warnung: Video konnte nicht eingebettet werden ({item.path}): {exc}")

    output_file.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_file)
    return stats


def write_run_report(path: Path, stats: dict, output_file: Path) -> None:
    report = {
        "output": str(output_file),
        "stats": stats,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "note": "Google Slides kann Video-Autoplay-Einstellungen aus PPTX ggf. nicht vollstaendig uebernehmen.",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, indent=2), encoding="utf-8")


def save_rendered_image(rendered: Image.Image, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    rendered.save(target, format="JPEG", quality=92)


def build_gslides_export_package(
    media_items: Iterable[MediaItem],
    package_dir: Path,
    ffmpeg_exe: Optional[str],
    conversion_cache_dir: Path,
) -> dict:
    package_dir.mkdir(parents=True, exist_ok=True)
    media_dir = package_dir / "media"
    media_dir.mkdir(parents=True, exist_ok=True)

    manifest_entries = []
    image_count = 0
    video_count = 0
    slide_index = 1

    for idx, item in enumerate(media_items, start=1):
        if item.kind == "image":
            img = load_image(item.path)
            rendered = render_stylized_canvas(img, item.width, item.height, is_low_quality(item))
            rel = Path("media") / f"{idx:04d}_image.jpg"
            save_rendered_image(rendered, package_dir / rel)
            manifest_entries.append(
                {
                    "slideIndex": slide_index,
                    "type": "image",
                    "captureTime": item.capture_time.isoformat(),
                    "sourcePath": str(item.path),
                    "file": str(rel).replace("\\", "/"),
                }
            )
            image_count += 1
            slide_index += 1
            continue

        if not ffmpeg_exe:
            print(f"Warnung: Video fuer GSlides-Paket uebersprungen (kein FFmpeg): {item.path}")
            continue

        try:
            poster_img = extract_video_frame(item.path, ffmpeg_exe)
            poster_canvas = render_stylized_canvas(
                poster_img, poster_img.width, poster_img.height, is_low_quality(item)
            )
            poster_rel = Path("media") / f"{idx:04d}_video_poster.jpg"
            save_rendered_image(poster_canvas, package_dir / poster_rel)

            embeddable = ensure_mp4_for_embedding(item.path, ffmpeg_exe, conversion_cache_dir)
            video_rel = Path("media") / f"{idx:04d}_video.mp4"
            shutil.copy2(embeddable, package_dir / video_rel)

            manifest_entries.append(
                {
                    "slideIndex": slide_index,
                    "type": "video_poster",
                    "captureTime": item.capture_time.isoformat(),
                    "sourcePath": str(item.path),
                    "file": str(poster_rel).replace("\\", "/"),
                }
            )
            slide_index += 1

            manifest_entries.append(
                {
                    "slideIndex": slide_index,
                    "type": "video",
                    "captureTime": item.capture_time.isoformat(),
                    "sourcePath": str(item.path),
                    "file": str(video_rel).replace("\\", "/"),
                    "autoplayRequested": True,
                    "mutedRequested": False,
                }
            )
            slide_index += 1
            video_count += 1
        except Exception as exc:
            print(f"Warnung: Video konnte nicht ins GSlides-Paket ({item.path}): {exc}")

    manifest = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "slideCountPlanned": len(manifest_entries),
        "images": image_count,
        "videos": video_count,
        "notes": [
            "Media-Dateien in einen Google-Drive-Ordner hochladen.",
            "Code.gs in Apps Script einfuegen und die Funktion buildPresentationFromManifest starten.",
            "Google Slides Autoplay ist von API-/Produktgrenzen abhaengig und ggf. nachzusetzen.",
        ],
        "entries": manifest_entries,
    }
    (package_dir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    apps_script_dir = package_dir / "apps_script"
    apps_script_dir.mkdir(parents=True, exist_ok=True)
    (apps_script_dir / "appsscript.json").write_text(
        json.dumps(
            {
                "timeZone": "Europe/Berlin",
                "dependencies": {},
                "exceptionLogging": "STACKDRIVER",
                "runtimeVersion": "V8",
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    (apps_script_dir / "Code.gs").write_text(_apps_script_code(), encoding="utf-8")

    return {"manifestEntries": len(manifest_entries), "images": image_count, "videos": video_count}


def _apps_script_code() -> str:
    return """function buildPresentationFromManifest() {
  // 1) Trage hier den Drive-Ordner ein, in dem manifest.json + media/ liegen.
  const folderId = 'REPLACE_WITH_DRIVE_FOLDER_ID';
  const folder = DriveApp.getFolderById(folderId);
  const manifestFile = getFileByName_(folder, 'manifest.json');
  if (!manifestFile) throw new Error('manifest.json nicht gefunden');
  const manifest = JSON.parse(manifestFile.getBlob().getDataAsString());

  const presentation = SlidesApp.create('Auto Generated Media Deck');
  const presId = presentation.getId();
  const slides = presentation.getSlides();
  if (slides.length > 0) {
    // Entferne initiale Leerseite.
    slides[0].remove();
  }

  manifest.entries.forEach((entry) => {
    const slide = presentation.appendSlide(SlidesApp.PredefinedLayout.BLANK);
    const file = getFileByPath_(folder, entry.file);
    if (!file) throw new Error('Datei nicht gefunden: ' + entry.file);

    if (entry.type === 'image' || entry.type === 'video_poster') {
      slide.insertImage(file.getBlob(), 0, 0, presentation.getPageWidth(), presentation.getPageHeight());
      return;
    }

    if (entry.type === 'video') {
      // Hinweis: SlidesApp erlaubt Drive-Videoeinbettung nicht in allen Konten/APIs gleich.
      // Wir versuchen zuerst den REST-Weg (createVideo + updateVideoProperties).
      createDriveVideoOnSlide_(presId, slide.getObjectId(), file.getId(), entry.autoplayRequested === true);
    }
  });

  Logger.log('Created presentation: https://docs.google.com/presentation/d/' + presId + '/edit');
}

function createDriveVideoOnSlide_(presentationId, slideObjectId, driveFileId, autoplayRequested) {
  // Aktivere in Apps Script: "Google Slides API" als Advanced Service.
  const requests = [
    {
      createVideo: {
        objectId: 'video_' + new Date().getTime(),
        elementProperties: {
          pageObjectId: slideObjectId,
          size: {
            width: { magnitude: 960, unit: 'PT' },
            height: { magnitude: 540, unit: 'PT' }
          },
          transform: {
            scaleX: 1,
            scaleY: 1,
            translateX: 0,
            translateY: 0,
            unit: 'PT'
          }
        },
        source: 'DRIVE',
        id: driveFileId
      }
    }
  ];

  const created = Slides.Presentations.batchUpdate({ requests: requests }, presentationId);
  const replies = created.replies || [];
  const createReply = replies.find(r => r.createVideo);
  if (!createReply) return;
  const videoObjectId = createReply.createVideo.objectId;
  if (!videoObjectId) return;

  if (autoplayRequested) {
    Slides.Presentations.batchUpdate(
      {
        requests: [
          {
            updateVideoProperties: {
              objectId: videoObjectId,
              videoProperties: {
                autoPlay: true
              },
              fields: 'autoPlay'
            }
          }
        ]
      },
      presentationId
    );
  }
}

function getFileByName_(folder, name) {
  const files = folder.getFilesByName(name);
  return files.hasNext() ? files.next() : null;
}

function getFileByPath_(rootFolder, relPath) {
  const parts = relPath.split('/');
  let currentFolder = rootFolder;
  for (let i = 0; i < parts.length - 1; i++) {
    const sub = currentFolder.getFoldersByName(parts[i]);
    if (!sub.hasNext()) return null;
    currentFolder = sub.next();
  }
  const files = currentFolder.getFilesByName(parts[parts.length - 1]);
  return files.hasNext() ? files.next() : null;
}
"""


def main() -> None:
    parser = argparse.ArgumentParser(description="Erstellt eine bild- und videobasierte PPTX (Google Slides kompatibel).")
    parser.add_argument("--input", default="images", help="Quellordner mit Medien (rekursiv).")
    parser.add_argument("--output", default="output/all_images.pptx", help="Ziel-PPTX fuer Vollmodus.")
    parser.add_argument("--test-output", default="output/test_preview.pptx", help="Ziel-PPTX fuer Testmodus.")
    parser.add_argument("--report", default="output/run_report.json", help="JSON-Bericht zur Generierung.")
    parser.add_argument("--test", action="store_true", help="Erstellt Test-PPTX mit 4 Bildtypen + 1 Video (2 Folien).")
    parser.add_argument(
        "--gslides-package-dir",
        default="output/gslides_package",
        help="Schreibt ein Google-Slides-Importpaket (manifest + media + apps_script).",
    )
    parser.add_argument(
        "--export-gslides-package",
        action="store_true",
        help="Erzeugt statt PPTX ein natives GSlides-Hilfspaket (Apps Script + Medien).",
    )
    args = parser.parse_args()

    input_dir = Path(args.input)
    if not input_dir.exists():
        raise RuntimeError(f"Input-Ordner existiert nicht: {input_dir}")

    ffmpeg_exe = resolve_ffmpeg_executable()
    all_items = build_media_items(input_dir, ffmpeg_exe)
    if not all_items:
        raise RuntimeError(f"Keine unterstuetzten Medien in {input_dir} gefunden.")

    conversion_cache_dir = Path("output") / "_converted_videos"
    if args.export_gslides_package:
        selected = pick_test_items(all_items, Path(args.test_output).parent) if args.test else all_items
        stats = build_gslides_export_package(selected, Path(args.gslides_package_dir), ffmpeg_exe, conversion_cache_dir)
        print(f"GSlides-Paket erstellt: {args.gslides_package_dir}")
        print(json.dumps(stats, indent=2))
        return

    if args.test:
        selected = pick_test_items(all_items, Path(args.test_output).parent)
        stats = create_presentation(selected, Path(args.test_output), ffmpeg_exe, conversion_cache_dir)
        write_run_report(Path(args.report), stats, Path(args.test_output))
        print(f"Test-Praesentation erstellt: {args.test_output}")
        print(json.dumps(stats, indent=2))
        return

    stats = create_presentation(all_items, Path(args.output), ffmpeg_exe, conversion_cache_dir)
    write_run_report(Path(args.report), stats, Path(args.output))
    print(f"Praesentation erstellt: {args.output}")
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()
